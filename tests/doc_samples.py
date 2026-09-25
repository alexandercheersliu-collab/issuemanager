"""测试样本构造：程序化生成文字版 PDF / 扫描版 PDF / DOCX。

不依赖外部文件，全部在 tmp_path 下即时生成；
PDF 中文用 pymupdf 内置 CJK 字体（china-s），无需系统字体。
"""
from __future__ import annotations

import io
from pathlib import Path

from PIL import Image, ImageDraw


def make_text_pdf(path: Path, pages: list[str]) -> Path:
    """文字版 PDF：每页一段文本（有文本层）。"""
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page(width=595, height=842)  # A4
        y = 72
        for line in text.split("\n"):
            page.insert_text((72, y), line, fontsize=12, fontname="china-s")
            y += 24
    doc.save(path)
    doc.close()
    return path


def make_scanned_pdf(path: Path, page_count: int = 1) -> Path:
    """扫描版 PDF：整页一张图片，无文本层。"""
    import pymupdf

    doc = pymupdf.open()
    for i in range(page_count):
        image = Image.new("RGB", (600, 800), (255, 255, 255))
        draw = ImageDraw.Draw(image)
        draw.rectangle([50, 50, 550, 200], outline=(0, 0, 0), width=2)
        draw.line([50, 300 + i * 10, 550, 300 + i * 10], fill=(0, 0, 0), width=1)
        stream = io.BytesIO()
        image.save(stream, format="PNG")

        page = doc.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=stream.getvalue())
    doc.save(path)
    doc.close()
    return path


def make_docx(path: Path, paragraphs: list[str], *, with_image: bool = True) -> Path:
    """DOCX：按顺序写段落；with_image 时在中间插入一张内嵌图片。"""
    import docx

    document = docx.Document()
    midpoint = max(len(paragraphs) // 2, 1)
    for index, text in enumerate(paragraphs):
        document.add_paragraph(text)
        if with_image and index == midpoint - 1:
            image = Image.new("RGB", (120, 60), (240, 240, 255))
            stream = io.BytesIO()
            image.save(stream, format="PNG")
            stream.seek(0)
            document.add_picture(stream)
    document.save(path)
    return path
